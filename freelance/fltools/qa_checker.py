"""⑧ セルフレビュー・品質判定チェッカー: 納品前のバグ・ミスを自動検出する。

ファイル(またはフォルダごと)を渡すと、形式ごとにチェックして Markdown レポートを出す。
ERROR が1件でもあれば終了コード1(=納品NG)。

共通: 仮置き文字(〇〇, XX, TODO, TBD, ダミー, {{ }}, 要確認…)、「。。」等の記号重複、
      表記ゆれ(下さい/ください 等)、半角カナ
xlsx: LibreOfficeで再計算した上での数式エラー(#REF!, #DIV/0!, #VALUE!, #NAME?, #N/A)、
      数値が文字列で入っているセル、空の見出し、非表示シート
docx: 空の見出し、見出しレベルの飛び、表の空セル
pptx: 1枚あたりの文字量過多、空スライド、スライド外にはみ出した図形
csv : 重複行、空欄率、前後空白、列数不一致
pdf : テキスト抽出可否(画像だけのPDFでないか)
"""

from __future__ import annotations

import re
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

PLACEHOLDER = re.compile(r"(〇〇|○○|◯◯|ＸＸ|XX+|xxx+|TODO|TBD|FIXME|ダミー|仮テキスト|要確認|要入力|lorem ipsum|\{\{.*?\}\}|＜.*?を入力＞|<.*?を入力>)", re.I)
DOUBLE_PUNCT = re.compile(r"(。。|、、|，，|．．|！！！+|\?\?|？？)")
HALF_KANA = re.compile(r"[ｦ-ﾟ]")
YURE = [("下さい", "ください"), ("出来る", "できる"), ("頂く", "いただく"), ("致します", "いたします"),
        ("御座います", "ございます"), ("及び", "および"), ("又は", "または"), ("事が", "ことが")]
XL_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")


@dataclass
class Report:
    file: str
    items: list[tuple[str, str, str]] = field(default_factory=list)  # (level, where, message)

    def add(self, level, where, msg):
        self.items.append((level, where, msg))

    @property
    def errors(self):
        return sum(1 for i in self.items if i[0] == "ERROR")


def check_text(rep: Report, where: str, text: str, seen_yure: dict) -> None:
    if not text:
        return
    for m in PLACEHOLDER.finditer(text):
        rep.add("ERROR", where, f"仮置き文字が残っています: 「{m.group(0)}」")
    for m in DOUBLE_PUNCT.finditer(text):
        rep.add("WARN", where, f"記号の重複: 「{m.group(0)}」")
    if HALF_KANA.search(text):
        rep.add("WARN", where, "半角カナが含まれています")
    for a, b in YURE:
        if a in text:
            seen_yure.setdefault((a, b), set()).add("a")
        if b in text:
            seen_yure.setdefault((a, b), set()).add("b")


def finish_yure(rep: Report, seen: dict) -> None:
    for (a, b), s in seen.items():
        if s == {"a", "b"}:
            rep.add("WARN", "全体", f"表記ゆれ: 「{a}」と「{b}」が混在しています")


def _recalc_xlsx(path: Path) -> Path | None:
    """LibreOfficeで開いて保存し直し、数式の計算結果をキャッシュさせる。"""
    from .pdf_tools import convert, soffice_bin

    if not soffice_bin():
        return None
    tmp = Path(tempfile.mkdtemp())
    src = tmp / ("src_" + path.name)
    shutil.copy(path, src)
    try:
        return convert(str(src), str(tmp / "out"), "xlsx")
    except Exception:
        return None


def check_xlsx(path: Path, rep: Report) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    recalced = _recalc_xlsx(path)
    wb_val = load_workbook(recalced, data_only=True) if recalced else None
    if not recalced:
        rep.add("INFO", "全体", "LibreOffice が無いため数式の再計算チェックは省略しました")
    seen: dict = {}
    formula_count = 0
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            rep.add("WARN", ws.title, "非表示シートがあります(納品物に不要なら削除)")
        wsv = wb_val[ws.title] if wb_val and ws.title in wb_val.sheetnames else None
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                where = f"{ws.title}!{c.coordinate}"
                if isinstance(v, str) and v.startswith("="):
                    formula_count += 1
                    if wsv is not None:
                        cv = wsv[c.coordinate].value
                        if isinstance(cv, str) and cv in XL_ERRORS:
                            rep.add("ERROR", where, f"数式エラー {cv}: {v}")
                elif isinstance(v, str):
                    if v in XL_ERRORS:
                        rep.add("ERROR", where, f"エラー値 {v}")
                    elif re.fullmatch(r"-?[\d,]+(\.\d+)?", v.strip()) and not v.startswith("0"):
                        rep.add("WARN", where, f"数値が文字列として入っています: {v!r}(集計されません)")
                    check_text(rep, where, v, seen)
        # 見出し行の空セル(オートフィルタ範囲があればその見出し行、なければ最初の2セル以上埋まった行)
        if ws.auto_filter.ref:
            from openpyxl.utils import range_boundaries

            c1, r1, c2, _ = range_boundaries(ws.auto_filter.ref)
            header = [ws.cell(row=r1, column=ci) for ci in range(c1, c2 + 1)]
        else:
            header = next((list(r) for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 5))
                           if sum(c.value is not None for c in r) >= 2), [])
            while header and header[-1].value is None:
                header.pop()
        for c in header:
            if c.value is None:
                rep.add("WARN", f"{ws.title}!{c.coordinate}", "見出し行に空セルがあります")
    rep.add("INFO", "全体", f"シート数 {len(wb.worksheets)} / 数式 {formula_count}個")
    finish_yure(rep, seen)


def check_docx(path: Path, rep: Report) -> None:
    from docx import Document

    doc = Document(path)
    seen: dict = {}
    prev_level = 0
    for i, p in enumerate(doc.paragraphs, 1):
        where = f"段落{i}"
        if p.style.name.startswith("Heading"):
            try:
                lvl = int(p.style.name.split()[-1])
            except ValueError:
                lvl = prev_level
            if not p.text.strip():
                rep.add("WARN", where, "空の見出しがあります")
            if prev_level and lvl > prev_level + 1:
                rep.add("WARN", where, f"見出しレベルが飛んでいます(H{prev_level}→H{lvl}): {p.text[:30]}")
            prev_level = lvl
        check_text(rep, where, p.text, seen)
    for ti, t in enumerate(doc.tables, 1):
        empty = 0
        for r in t.rows:
            for c in r.cells:
                check_text(rep, f"表{ti}", c.text, seen)
                empty += not c.text.strip()
        if empty:
            rep.add("WARN", f"表{ti}", f"空セルが{empty}個あります")
    rep.add("INFO", "全体", f"段落 {len(doc.paragraphs)} / 表 {len(doc.tables)}")
    finish_yure(rep, seen)


def check_pptx(path: Path, rep: Report, max_chars: int = 250) -> None:
    from pptx import Presentation

    prs = Presentation(path)
    seen: dict = {}
    W, H = prs.slide_width, prs.slide_height
    for i, s in enumerate(prs.slides, 1):
        where = f"スライド{i}"
        chars = 0
        for shp in s.shapes:
            if shp.left is not None and (shp.left + shp.width > W * 1.01 or shp.top + shp.height > H * 1.01 or shp.left < 0 or shp.top < 0):
                rep.add("ERROR", where, f"図形がスライドの外にはみ出しています: {shp.name}")
            texts = []
            if shp.has_text_frame:
                texts.append(shp.text_frame.text)
            if getattr(shp, "has_table", False) and shp.has_table:
                texts += [c.text for r in shp.table.rows for c in r.cells]
            for t in texts:
                chars += len(t)
                check_text(rep, where, t, seen)
        if chars == 0 and not any(sh.shape_type == 13 for sh in s.shapes):
            rep.add("WARN", where, "文字も画像も無いスライドです")
        if chars > max_chars:
            rep.add("WARN", where, f"文字量が多すぎます({chars}字 > {max_chars}字)。分割を検討")
    rep.add("INFO", "全体", f"スライド {len(prs.slides)}枚")
    finish_yure(rep, seen)


def check_csv(path: Path, rep: Report) -> None:
    import csv

    from .common import detect_encoding

    enc = detect_encoding(path)
    if enc not in ("utf-8-sig",):
        rep.add("INFO", "全体", f"文字コード {enc}(Excelで開く納品なら UTF-8 BOM付き推奨)")
    with open(path, encoding=enc, newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        rep.add("ERROR", "全体", "空のファイルです")
        return
    header, body = rows[0], rows[1:]
    seen: dict = {}
    if len(set(header)) != len(header):
        rep.add("ERROR", "1行目", "列名が重複しています")
    dup = len(body) - len({tuple(r) for r in body})
    if dup:
        rep.add("WARN", "全体", f"重複行が{dup}行あります")
    for i, r in enumerate(body, 2):
        if len(r) != len(header):
            rep.add("ERROR", f"{i}行目", f"列数が見出しと違います({len(r)} ≠ {len(header)})")
        for j, v in enumerate(r):
            if v != v.strip():
                rep.add("WARN", f"{i}行目 {header[j] if j < len(header) else j}", "前後に空白があります")
            check_text(rep, f"{i}行目", v, seen)
    for j, h in enumerate(header):
        blanks = sum(1 for r in body if j >= len(r) or not r[j].strip())
        if body and blanks / len(body) > 0.3:
            rep.add("WARN", f"列「{h}」", f"空欄率 {blanks / len(body):.0%}")
    rep.add("INFO", "全体", f"{len(body)}行 × {len(header)}列")
    finish_yure(rep, seen)


def check_pdf(path: Path, rep: Report) -> None:
    from pypdf import PdfReader

    r = PdfReader(path)
    seen: dict = {}
    empty = 0
    for i, p in enumerate(r.pages, 1):
        t = p.extract_text() or ""
        empty += not t.strip()
        check_text(rep, f"p.{i}", t, seen)
    if empty == len(r.pages):
        rep.add("WARN", "全体", "テキストが抽出できません(画像のみのPDF?検索・コピー不可)")
    rep.add("INFO", "全体", f"{len(r.pages)}ページ")
    finish_yure(rep, seen)


def check_textfile(path: Path, rep: Report) -> None:
    seen: dict = {}
    for i, line in enumerate(path.read_text(encoding="utf-8", errors="replace").splitlines(), 1):
        check_text(rep, f"{i}行目", line, seen)
    finish_yure(rep, seen)


CHECKERS = {".xlsx": check_xlsx, ".xlsm": check_xlsx, ".docx": check_docx, ".pptx": check_pptx,
            ".csv": check_csv, ".pdf": check_pdf, ".md": check_textfile, ".txt": check_textfile}


def check(path: Path) -> Report:
    rep = Report(str(path))
    fn = CHECKERS.get(path.suffix.lower())
    if not fn:
        rep.add("INFO", "全体", "対象外の形式のためスキップ")
        return rep
    try:
        fn(path, rep)
    except Exception as e:  # 壊れたファイル自体が重大な欠陥
        rep.add("ERROR", "全体", f"ファイルを開けません: {e}")
    return rep


def render(reports: list[Report], limit: int = 50) -> str:
    total_err = sum(r.errors for r in reports)
    total_warn = sum(1 for r in reports for i in r.items if i[0] == "WARN")
    verdict = "❌ 納品NG(ERRORを修正してください)" if total_err else ("⚠️ 条件付きOK(WARNを確認)" if total_warn else "✅ 納品OK")
    out = [f"# 品質チェックレポート", "", f"**判定: {verdict}**  (ERROR {total_err} / WARN {total_warn})", ""]
    for r in reports:
        out.append(f"## {r.file}")
        if not r.items:
            out.append("- 問題なし")
        shown = sorted(r.items, key=lambda x: {"ERROR": 0, "WARN": 1, "INFO": 2}[x[0]])
        for level, where, msg in shown[:limit]:
            out.append(f"- **{level}** [{where}] {msg}")
        if len(shown) > limit:
            out.append(f"- …ほか {len(shown) - limit}件")
        out.append("")
    return "\n".join(out)


def run(args) -> int:
    targets: list[Path] = []
    for t in args.targets:
        p = Path(t)
        targets += sorted(x for x in p.rglob("*") if x.suffix.lower() in CHECKERS) if p.is_dir() else [p]
    reports = [check(p) for p in targets]
    text = render(reports)
    print(text)
    if args.out:
        Path(args.out).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out).write_text(text, encoding="utf-8")
    return 1 if any(r.errors for r in reports) else 0


def register(sub) -> None:
    p = sub.add_parser("check", help="⑧ 納品前の品質チェック(ファイル or フォルダ)")
    p.add_argument("targets", nargs="+")
    p.add_argument("-o", "--out", help="レポートの保存先(.md)")
    p.set_defaults(func=run)
