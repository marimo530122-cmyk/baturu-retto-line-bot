"""② xlsx自動構築エンジン: JSON仕様 or CSV から、関数・書式・レイアウト込みの Excel を生成する。

仕様(JSON)の例は templates/specs/sample_xlsx.json を参照。

- 列定義 `columns` の `formula` には `{row}`(行番号)と `{col:キー}`(その列の列記号)を使える。
  例: "=({col:qty}{row}*{col:price}{row})"
- `totals` で列ごとに SUM/AVERAGE/COUNT/MAX/MIN の集計行を自動追加(関数で入る=後から編集しても追従)。
- 値はハードコードせず、計算は必ず Excel 関数として入れる(納品後にクライアントが数値を変えても壊れない)。
"""

from __future__ import annotations

import csv
import datetime
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .common import JP_FONT, detect_encoding, ensure_parent, load_spec

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=JP_FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=JP_FONT)
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
BAND_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _east_asian_width(s: str) -> int:
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(s))


def _coerce(value):
    """CSV 由来の文字列を数値・日付に変換できるなら変換する(先頭0のコード類は文字列のまま)。"""
    if not isinstance(value, str):
        return value
    m = re.fullmatch(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", value.strip())
    if m:
        try:
            return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return value
    v = value.strip().replace(",", "")
    if v == "" or (len(v) > 1 and v.startswith("0") and not v.startswith("0.")):
        return value
    try:
        return int(v)
    except ValueError:
        try:
            return float(v)
        except ValueError:
            return value


def build_sheet(wb: Workbook, sheet_spec: dict, first: bool) -> None:
    ws = wb.active if first else wb.create_sheet()
    ws.title = sheet_spec.get("name", "Sheet1")[:31]
    columns = sheet_spec.get("columns") or []
    rows = sheet_spec.get("rows") or []

    if not columns and rows and isinstance(rows[0], dict):
        columns = [{"header": k, "key": k} for k in rows[0].keys()]
    keys = [c.get("key", c["header"]) for c in columns]
    letter_of = {k: get_column_letter(i + 1) for i, k in enumerate(keys)}
    letter_of.update({c["header"]: letter_of[k] for c, k in zip(columns, keys)})

    r = 1
    if sheet_spec.get("title"):
        ws.cell(row=1, column=1, value=sheet_spec["title"]).font = Font(name=JP_FONT, bold=True, size=14)
        if len(columns) > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        r = 3
    header_row = r
    for ci, col in enumerate(columns, 1):
        c = ws.cell(row=header_row, column=ci, value=col["header"])
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    first_data = header_row + 1
    for ri, row in enumerate(rows):
        excel_row = first_data + ri
        for ci, col in enumerate(columns, 1):
            key = keys[ci - 1]
            if col.get("formula"):
                value = _expand(col["formula"], excel_row, letter_of)
            elif isinstance(row, dict):
                value = _coerce(row.get(key, row.get(col["header"])))
            else:
                value = _coerce(row[ci - 1]) if ci - 1 < len(row) else None
            c = ws.cell(row=excel_row, column=ci, value=value)
            c.font, c.border = BODY_FONT, BORDER
            if col.get("format"):
                c.number_format = col["format"]
            elif isinstance(value, datetime.date):
                c.number_format = "yyyy/mm/dd"
            if col.get("align"):
                c.alignment = Alignment(horizontal=col["align"])
            if sheet_spec.get("banded", True) and ri % 2 == 1:
                c.fill = BAND_FILL
    last_data = first_data + len(rows) - 1

    totals = sheet_spec.get("totals") or {}
    if totals and rows:
        tr = last_data + 1
        label_col = sheet_spec.get("totals_label_col", 1)
        ws.cell(row=tr, column=label_col, value=sheet_spec.get("totals_label", "合計"))
        for ci, col in enumerate(columns, 1):
            func = totals.get(keys[ci - 1]) or totals.get(col["header"])
            if func:
                L = get_column_letter(ci)
                c = ws.cell(row=tr, column=ci, value=f"={func.upper()}({L}{first_data}:{L}{last_data})")
                if col.get("format"):
                    c.number_format = col["format"]
        for ci in range(1, len(columns) + 1):
            c = ws.cell(row=tr, column=ci)
            c.fill, c.border = TOTAL_FILL, BORDER
            c.font = Font(name=JP_FONT, bold=True)

    # 列幅: 指定があればそれ、なければ中身から自動計算(全角=2)
    for ci, col in enumerate(columns, 1):
        if col.get("width"):
            width = col["width"]
        else:
            sample = [col["header"]] + [
                (r_.get(keys[ci - 1], "") if isinstance(r_, dict) else (r_[ci - 1] if ci - 1 < len(r_) else ""))
                for r_ in rows[:200]
            ]
            width = min(max(_east_asian_width(str(s)) for s in sample) + 2, 60)
            width = max(width, 8)
        ws.column_dimensions[get_column_letter(ci)].width = width

    if sheet_spec.get("freeze", True):
        ws.freeze_panes = ws.cell(row=first_data, column=sheet_spec.get("freeze_cols", 0) + 1)
    if sheet_spec.get("autofilter", True) and columns:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{max(last_data, header_row)}"

    # 入力規則(プルダウン)
    for ci, col in enumerate(columns, 1):
        if col.get("choices"):
            L = get_column_letter(ci)
            dv = DataValidation(type="list", formula1='"' + ",".join(col["choices"]) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{L}{first_data}:{L}{max(last_data, first_data) + sheet_spec.get('spare_rows', 100)}")

    # 条件付き書式: {"column": "key", "operator": "lessThan", "value": 0, "color": "FFC7CE"}
    for rule in sheet_spec.get("conditional", []):
        L = letter_of[rule["column"]]
        ws.conditional_formatting.add(
            f"{L}{first_data}:{L}{max(last_data, first_data)}",
            CellIsRule(operator=rule["operator"], formula=[str(rule["value"])],
                       fill=PatternFill("solid", fgColor=rule.get("color", "FFC7CE"))),
        )

    # 自由配置セル(集計ボックス等): {"cell": "H3", "value": "=SUM(...)", "format": "#,##0", "bold": true}
    for extra in sheet_spec.get("cells", []):
        c = ws[extra["cell"]]
        c.value = _expand(extra["value"], 0, letter_of) if isinstance(extra["value"], str) else extra["value"]
        c.font = Font(name=JP_FONT, bold=extra.get("bold", False))
        if extra.get("format"):
            c.number_format = extra["format"]

    ws.sheet_view.zoomScale = sheet_spec.get("zoom", 100)
    ws.page_setup.orientation = sheet_spec.get("orientation", "landscape")
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{header_row}:{header_row}"


def _expand(formula: str, row: int, letter_of: dict) -> str:
    out = formula.replace("{row}", str(row))
    for k, L in letter_of.items():
        out = out.replace("{col:" + k + "}", L)
    return out


def build_from_spec(spec: dict, out: str) -> Path:
    wb = Workbook()
    for i, sheet in enumerate(spec["sheets"]):
        build_sheet(wb, sheet, first=(i == 0))
    path = ensure_parent(out)
    wb.save(path)
    return path


def build_from_csv(csv_path: str, out: str, sheet_name: str = "データ", title: str | None = None) -> Path:
    enc = detect_encoding(csv_path)
    with open(csv_path, encoding=enc, newline="") as f:
        data = list(csv.reader(f))
    header, body = data[0], data[1:]
    spec = {"sheets": [{"name": sheet_name, "title": title,
                        "columns": [{"header": h, "key": h} for h in header], "rows": body}]}
    return build_from_spec(spec, out)


def run(args) -> int:
    if args.spec:
        path = build_from_spec(load_spec(args.spec), args.out)
    elif args.csv:
        path = build_from_csv(args.csv, args.out, title=args.title)
    else:
        raise SystemExit("--spec か --csv のどちらかを指定してください")
    print(f"[xlsx] 生成しました: {path}")
    return 0


def register(sub) -> None:
    p = sub.add_parser("xlsx", help="② Excel を JSON仕様 or CSV から自動構築")
    p.add_argument("--spec", help="JSON仕様ファイル")
    p.add_argument("--csv", help="CSVからそのまま整形済みExcelを作る")
    p.add_argument("--title", help="CSVモード時のシートタイトル")
    p.add_argument("-o", "--out", required=True, help="出力 .xlsx")
    p.set_defaults(func=run)
