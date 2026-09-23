"""④ pptx自動生成モジュール: Markdown からプレゼン資料(16:9)を構築する。

Markdown の書き方:
  # 資料タイトル            ← 表紙スライド(直後の段落はサブタイトル)
  ## スライド見出し          ← 1枚のスライド
  - 箇条書き(インデントで2階層まで)
  | 表 | も | 書ける |      ← 表スライド
  > 発表者ノートに入れたい文章
  ### セクション名           ← 中扉(セクション区切り)スライド

文字量に応じてフォントサイズを自動縮小し、はみ出しを防ぐ。
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .common import JP_FONT, ensure_parent, parse_markdown, split_bold, strip_inline_md

W, H = Inches(13.333), Inches(7.5)
ACCENT = RGBColor(0x1F, 0x4E, 0x78)
SUB = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _font(run, size, bold=False, color=None):
    run.font.name = JP_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = rpr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}ea", {})
        rpr.append(ea)
    ea.set("typeface", JP_FONT)
    if color is not None:
        run.font.color.rgb = color


def _textbox(slide, left, top, width, height, text, size, bold=False, color=None, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _font(run, size, bold, color)
    return tb


def _rect(slide, left, top, width, height, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    return shp


def _slide_number(slide, n):
    _textbox(slide, W - Inches(1.2), H - Inches(0.5), Inches(1), Inches(0.35), str(n), 10, color=SUB, align=PP_ALIGN.RIGHT)


def _title_slide(prs, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, W, H, ACCENT)
    _textbox(s, Inches(0.8), Inches(2.4), W - Inches(1.6), Inches(1.5), title, 40, True, WHITE)
    if subtitle:
        _textbox(s, Inches(0.8), Inches(4.0), W - Inches(1.6), Inches(1), subtitle, 20, color=WHITE)
    return s


def _section_slide(prs, title, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, Inches(3.0), W, Inches(1.5), ACCENT)
    _textbox(s, Inches(0.8), Inches(3.25), W - Inches(1.6), Inches(1), title, 32, True, WHITE)
    _slide_number(s, n)
    return s


def _content_slide(prs, title, blocks, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, W, Inches(1.1), ACCENT)
    _textbox(s, Inches(0.6), Inches(0.22), W - Inches(1.2), Inches(0.8), strip_inline_md(title), 28, True, WHITE)
    _slide_number(s, n)

    top = Inches(1.4)
    body_h = H - top - Inches(0.7)
    tables = [b for b in blocks if b["type"] == "table"]
    texts = [b for b in blocks if b["type"] != "table"]

    lines = []  # (text, level, bullet?)
    for b in texts:
        if b["type"] == "paragraph":
            lines.append((b["text"], 0, False))
        elif b["type"] in ("bullets", "numbered"):
            for i, it in enumerate(b["items"], 1):
                prefix = f"{i}. " if b["type"] == "numbered" and it["level"] == 0 else ("・" if it["level"] == 0 else "－ ")
                lines.append((prefix + it["text"], it["level"], True))
        elif b["type"] == "heading":
            lines.append((f"**{b['text']}**", 0, False))

    text_h = body_h if not tables else (body_h * 0.35 if lines else 0)
    if lines:
        total_chars = sum(len(strip_inline_md(t)) for t, _, _ in lines)
        size = 24 if len(lines) <= 5 and total_chars < 150 else 20 if len(lines) <= 8 and total_chars < 300 else 16 if total_chars < 500 else 13
        if tables:
            size = min(size, 16)
        tb = s.shapes.add_textbox(Inches(0.8), top, W - Inches(1.6), text_h)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, (text, level, _) in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.level = level
            p.space_after = Pt(size * 0.5)
            for chunk, bold in split_bold(text):
                r = p.add_run()
                r.text = chunk
                _font(r, size - 2 * level, bold, ACCENT if bold else None)

    for tbl in tables[:1]:
        headers, rows = tbl["headers"], tbl["rows"]
        t_top = top + (text_h if lines else 0)
        t_h = min(Inches(0.45) * (len(rows) + 1), H - t_top - Inches(0.7))
        shape = s.shapes.add_table(len(rows) + 1, len(headers), Inches(0.8), t_top, W - Inches(1.6), Emu(int(t_h)))
        size = 14 if len(rows) <= 8 else 11 if len(rows) <= 14 else 9
        for ci, h in enumerate(headers):
            cell = shape.table.cell(0, ci)
            cell.text = strip_inline_md(h)
            _font(cell.text_frame.paragraphs[0].runs[0], size, True, WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT
        for ri, row in enumerate(rows, 1):
            for ci in range(len(headers)):
                cell = shape.table.cell(ri, ci)
                cell.text = strip_inline_md(row[ci] if ci < len(row) else "")
                if cell.text_frame.paragraphs[0].runs:
                    _font(cell.text_frame.paragraphs[0].runs[0], size)
    return s


def build(markdown: str, out: str) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # スライド単位に分割(ノートは "> " 行)
    chunks: list[dict] = []
    cur = None
    for line in markdown.splitlines():
        st = line.strip()
        if st.startswith("# ") and not chunks and cur is None:
            cur = {"kind": "title", "title": st[2:].strip(), "body": [], "notes": []}
            chunks.append(cur)
        elif st.startswith("## "):
            cur = {"kind": "content", "title": st[3:].strip(), "body": [], "notes": []}
            chunks.append(cur)
        elif st.startswith("### "):
            cur = {"kind": "section", "title": st[4:].strip(), "body": [], "notes": []}
            chunks.append(cur)
        elif cur is not None:
            if st.startswith(">"):
                cur["notes"].append(st.lstrip("> ").strip())
            else:
                cur["body"].append(line)

    n = 0
    for c in chunks:
        n += 1
        if c["kind"] == "title":
            sub = " ".join(l.strip() for l in c["body"] if l.strip())
            slide = _title_slide(prs, strip_inline_md(c["title"]), strip_inline_md(sub))
        elif c["kind"] == "section":
            slide = _section_slide(prs, strip_inline_md(c["title"]), n)
        else:
            slide = _content_slide(prs, c["title"], parse_markdown("\n".join(c["body"])), n)
        if c["notes"]:
            slide.notes_slide.notes_text_frame.text = "\n".join(c["notes"])

    path = ensure_parent(out)
    prs.save(path)
    return path


def run(args) -> int:
    path = build(Path(args.md).read_text(encoding="utf-8"), args.out)
    print(f"[pptx] 生成しました: {path}")
    return 0


def register(sub) -> None:
    p = sub.add_parser("pptx", help="④ PowerPoint を Markdown から構築")
    p.add_argument("--md", required=True, help="構成 Markdown ファイル")
    p.add_argument("-o", "--out", required=True)
    p.set_defaults(func=run)
