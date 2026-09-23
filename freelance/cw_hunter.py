"""クラウドワークス案件自動ハンター

公開の案件検索ページ(ログイン不要)を指定キーワードで検索し、
「AIで9割自動化できて安全な事務系案件」だけを点数付きで一覧にする。

- 取得: 検索結果ページに埋め込まれた案件データ(JSON)を読む。
  読めなかったときは、ページ内の案件リンクだけを拾う簡易モードに落ちる。
- 判定: config.json の NG ワード・怪しいワード・最低報酬で除外し、
  対象ワードの一致と報酬額で点数を付ける。
- 出力: freelance/projects/ に CSV(Excelでそのまま開ける UTF-8 BOM 付き)。
  openpyxl が入っていれば同じ内容の .xlsx も作る。

サイトに負荷をかけないよう、robots.txt で禁止されていれば取得せず、
リクエストの間隔は config.json の request_interval_sec 秒以上あける。
"""

import csv
import datetime as dt
import html
import json
import re
import time
import urllib.error
import urllib.parse
import urllib.request
import urllib.robotparser
from html.parser import HTMLParser
from pathlib import Path

BASE_URL = "https://crowdworks.jp"
SEARCH_PATH = "/public/jobs/search"
USER_AGENT = "Mozilla/5.0 (compatible; cw-hunter/1.0; personal job patrol)"

FREELANCE_DIR = Path(__file__).resolve().parent
PROJECTS_DIR = FREELANCE_DIR / "projects"
SEEN_PATH = PROJECTS_DIR / "seen_ids.json"
CONFIG_PATH = FREELANCE_DIR / "config.json"

COLUMNS = [
    ("new", "新着"),
    ("score", "スコア"),
    ("title", "タイトル"),
    ("budget_text", "報酬"),
    ("payment_type", "形式"),
    ("url", "URL"),
    ("summary", "概要"),
    ("reasons", "判定理由"),
    ("matched_keyword", "検索キーワード"),
    ("id", "案件ID"),
]


class FetchBlocked(Exception):
    """robots.txt で禁止されている、またはサイトに拒否されたので取得を止める。"""


# ---------------------------------------------------------------- 取得


def search_url(keyword, page=1):
    params = {
        "search[keywords]": keyword,
        "order": "new",
        "hide_expired": "true",
    }
    if page > 1:
        params["page"] = str(page)
    return BASE_URL + SEARCH_PATH + "?" + urllib.parse.urlencode(params)


def load_robots():
    robots = urllib.robotparser.RobotFileParser(BASE_URL + "/robots.txt")
    try:
        robots.read()
    except (urllib.error.URLError, OSError) as err:
        raise FetchBlocked(f"robots.txt を取得できませんでした: {err}") from err
    return robots


def fetch(url, robots):
    if not robots.can_fetch(USER_AGENT, url):
        raise FetchBlocked(f"robots.txt で取得が禁止されています: {url}")
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    try:
        with urllib.request.urlopen(request, timeout=30) as res:
            return res.read().decode("utf-8", errors="replace")
    except urllib.error.HTTPError as err:
        if err.code in (403, 429):
            raise FetchBlocked(f"サイトに拒否されました (HTTP {err.code})。しばらく時間をおいてください") from err
        raise


# ---------------------------------------------------------------- 解析


class _JsonAttrCollector(HTMLParser):
    """タグの属性や <script type="application/json"> に入った JSON をすべて集める。"""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blobs = []
        self.links = []
        self._in_json_script = False
        self._in_link = None
        self._link_text = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        for value in attrs.values():
            if value and value.lstrip()[:1] in "{[":
                self.blobs.append(value)
        if tag == "script" and "json" in (attrs.get("type") or ""):
            self._in_json_script = True
        if tag == "a":
            match = re.search(r"/public/jobs/(\d+)(?:[?#/]|$)", attrs.get("href") or "")
            if match:
                self._in_link = match.group(1)
                self._link_text = []

    def handle_endtag(self, tag):
        if tag == "script":
            self._in_json_script = False
        if tag == "a" and self._in_link:
            text = re.sub(r"\s+", " ", "".join(self._link_text)).strip()
            if text:
                self.links.append((self._in_link, text))
            self._in_link = None

    def handle_data(self, data):
        if self._in_json_script and data.strip():
            self.blobs.append(data)
        if self._in_link:
            self._link_text.append(data)


def _walk(node):
    if isinstance(node, dict):
        yield node
        for value in node.values():
            yield from _walk(value)
    elif isinstance(node, list):
        for value in node:
            yield from _walk(value)


def _numbers_by_key(node, predicate):
    """node 以下で、キー名が predicate を満たす数値をすべて集める。"""
    found = []
    for d in _walk(node):
        for key, value in d.items():
            if isinstance(value, (int, float)) and not isinstance(value, bool) and predicate(key):
                found.append((key, value))
    return found


def _payment_info(entry):
    """案件1件ぶんのデータから、報酬形式と金額の範囲を推定する。"""
    payment = entry.get("payment") if isinstance(entry.get("payment"), dict) else entry
    keys = " ".join(k for d in _walk(payment) for k in d.keys())

    if "hourly" in keys:
        values = [v for k, v in _numbers_by_key(payment, lambda k: "wage" in k or "hourly" in k) if v > 0]
        kind = "時間単価"
    elif "task" in keys:
        values = [v for k, v in _numbers_by_key(payment, lambda k: "price" in k or "reward" in k) if v > 0]
        kind = "タスク"
    else:
        values = [v for k, v in _numbers_by_key(payment, lambda k: "budget" in k or "price" in k) if v > 0]
        kind = "固定報酬"

    if not values:
        return {"payment_type": kind, "budget_min": None, "budget_max": None}
    return {"payment_type": kind, "budget_min": int(min(values)), "budget_max": int(max(values))}


def _job_from_entry(entry):
    offer = entry.get("job_offer") if isinstance(entry.get("job_offer"), dict) else entry
    job_id = offer.get("id")
    title = offer.get("title")
    if not job_id or not isinstance(title, str):
        return None
    summary = offer.get("description_digest") or offer.get("description") or ""
    job = {
        "id": str(job_id),
        "title": title.strip(),
        "summary": re.sub(r"\s+", " ", html.unescape(str(summary))).strip(),
        "url": f"{BASE_URL}/public/jobs/{job_id}",
    }
    job.update(_payment_info(entry))
    return job


def parse_search_page(page_html):
    """検索結果ページの HTML から案件のリストを返す。

    まず埋め込み JSON の中から job_offer(id と title を持つ)を探し、
    1件も見つからなければ案件ページへのリンクだけを拾う。
    """
    collector = _JsonAttrCollector()
    collector.feed(page_html)

    jobs = {}
    for blob in collector.blobs:
        try:
            data = json.loads(blob)
        except ValueError:
            continue
        for d in _walk(data):
            offer = d.get("job_offer")
            if isinstance(offer, dict) and "id" in offer and "title" in offer:
                job = _job_from_entry(d)
                if job:
                    jobs.setdefault(job["id"], job)

    if jobs:
        return list(jobs.values())

    for job_id, text in collector.links:
        jobs.setdefault(job_id, {
            "id": job_id,
            "title": text,
            "summary": "",
            "url": f"{BASE_URL}/public/jobs/{job_id}",
            "payment_type": "不明",
            "budget_min": None,
            "budget_max": None,
        })
    return list(jobs.values())


# ---------------------------------------------------------------- 判定


def _hits(text, words):
    lowered = text.lower()
    return [w for w in words if w.lower() in lowered]


def evaluate(job, config):
    """除外理由があれば excluded=True。なければ点数と理由を付ける。"""
    text = f"{job['title']} {job['summary']}"
    reasons = []
    excluded = False

    for word in _hits(text, config["ng_words"]):
        reasons.append(f"除外: {config['ng_words'][word]}({word})")
        excluded = True
    for word in _hits(text, config["scam_words"]):
        reasons.append(f"除外: 要注意案件 - {config['scam_words'][word]}({word})")
        excluded = True

    budget = job["budget_max"]
    kind = job["payment_type"]
    if budget is None:
        reasons.append("報酬不明(要確認)")
    elif kind == "時間単価" and budget < config["min_hourly_wage_yen"]:
        reasons.append(f"除外: 時給が低い({budget}円)")
        excluded = True
    elif kind != "時間単価" and budget < config["min_fixed_budget_yen"]:
        reasons.append(f"除外: 報酬が低い({budget}円)")
        excluded = True

    score = 0
    title_hits = _hits(job["title"], config["target_words"])
    body_hits = [w for w in _hits(job["summary"], config["target_words"]) if w not in title_hits]
    score += min(8, sum(config["target_words"][w] * 2 for w in title_hits)
                 + sum(config["target_words"][w] for w in body_hits))
    if title_hits or body_hits:
        reasons.append("対象: " + "・".join(title_hits + body_hits))

    if budget is not None:
        if kind == "時間単価":
            score += 2 if budget >= 1500 else 1 if budget >= 1000 else 0
        else:
            score += 3 if budget >= 30000 else 2 if budget >= 10000 else 1 if budget >= 5000 else 0

    if _hits(text, ["本日中", "即日納品", "大至急"]):
        score -= 1
        reasons.append("短納期(-1)")

    job["score"] = score
    job["excluded"] = excluded
    job["reasons"] = " / ".join(reasons)
    job["budget_text"] = budget_text(job)
    return job


def budget_text(job):
    lo, hi = job["budget_min"], job["budget_max"]
    if hi is None:
        return "不明"
    unit = "円/時" if job["payment_type"] == "時間単価" else "円"
    if lo == hi:
        return f"{hi:,}{unit}"
    return f"{lo:,}〜{hi:,}{unit}"


# ---------------------------------------------------------------- 出力


def load_config():
    return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))


def load_seen():
    try:
        return set(json.loads(SEEN_PATH.read_text(encoding="utf-8")))
    except (OSError, ValueError):
        return set()


def save_seen(ids):
    SEEN_PATH.write_text(json.dumps(sorted(ids), ensure_ascii=False), encoding="utf-8")


def write_csv(path, jobs):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([label for _, label in COLUMNS])
        for job in jobs:
            writer.writerow([job.get(key, "") for key, _ in COLUMNS])


def write_xlsx(path, selected, excluded):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return False

    wb = Workbook()
    for sheet_index, (title, jobs) in enumerate([("厳選案件", selected), ("除外した案件", excluded)]):
        ws = wb.active if sheet_index == 0 else wb.create_sheet()
        ws.title = title
        ws.append([label for _, label in COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
        for job in jobs:
            ws.append([job.get(key, "") for key, _ in COLUMNS])
            url_cell = ws.cell(row=ws.max_row, column=[k for k, _ in COLUMNS].index("url") + 1)
            url_cell.hyperlink = url_cell.value
            url_cell.font = Font(color="0563C1", underline="single")
        widths = {"new": 6, "score": 7, "title": 45, "budget_text": 18, "payment_type": 10,
                  "url": 40, "summary": 70, "reasons": 40, "matched_keyword": 14, "id": 10}
        for i, (key, _) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths[key]
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return True


# ---------------------------------------------------------------- 実行


def collect(keywords, config, from_html=None, log=print):
    """キーワードごとに検索して案件を集める。from_html を渡すと保存済み HTML を読む。"""
    jobs = {}

    def add(found, keyword):
        for job in found:
            job.setdefault("matched_keyword", keyword)
            jobs.setdefault(job["id"], job)

    if from_html:
        for path in from_html:
            found = parse_search_page(Path(path).read_text(encoding="utf-8", errors="replace"))
            log(f"  {path}: {len(found)}件")
            add(found, Path(path).stem)
        return list(jobs.values())

    robots = load_robots()
    first = True
    for keyword in keywords:
        for page in range(1, config["max_pages_per_keyword"] + 1):
            if not first:
                time.sleep(config["request_interval_sec"])
            first = False
            url = search_url(keyword, page)
            page_html = fetch(url, robots)
            found = parse_search_page(page_html)
            log(f"  「{keyword}」{page}ページ目: {len(found)}件")
            if not found:
                debug_path = PROJECTS_DIR / "debug" / f"search_{keyword}_{page}.html"
                debug_path.parent.mkdir(parents=True, exist_ok=True)
                debug_path.write_text(page_html, encoding="utf-8")
                log(f"    案件を読み取れませんでした。ページを {debug_path} に保存しました")
                break
            add(found, keyword)
    return list(jobs.values())


def run(keywords=None, from_html=None, min_budget=None, only_new=False, log=print):
    config = load_config()
    if min_budget is not None:
        config["min_fixed_budget_yen"] = min_budget
    keywords = keywords or config["search_keywords"]

    PROJECTS_DIR.mkdir(parents=True, exist_ok=True)
    log("案件を集めています…")
    jobs = collect(keywords, config, from_html=from_html, log=log)

    seen = load_seen()
    for job in jobs:
        evaluate(job, config)
        job["new"] = "" if job["id"] in seen else "★"

    selected = [j for j in jobs if not j["excluded"] and j["score"] >= config["min_score"]]
    excluded = [j for j in jobs if j["excluded"] or j["score"] < config["min_score"]]
    for j in excluded:
        if not j["excluded"]:
            j["reasons"] = (j["reasons"] + " / " if j["reasons"] else "") + f"除外: スコア不足({j['score']})"
    if only_new:
        selected = [j for j in selected if j["new"]]
    selected.sort(key=lambda j: (-j["score"], j["new"] != "★", -(j["budget_max"] or 0)))

    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M")
    csv_path = PROJECTS_DIR / f"cw_{stamp}.csv"
    excluded_path = PROJECTS_DIR / f"cw_{stamp}_excluded.csv"
    write_csv(csv_path, selected)
    write_csv(excluded_path, excluded)
    xlsx_path = PROJECTS_DIR / f"cw_{stamp}.xlsx"
    wrote_xlsx = write_xlsx(xlsx_path, selected, excluded)

    if not from_html:
        save_seen(seen | {j["id"] for j in jobs})

    new_count = sum(1 for j in selected if j["new"])
    log("")
    log(f"取得 {len(jobs)}件 → 厳選 {len(selected)}件(うち新着 {new_count}件) / 除外 {len(excluded)}件")
    for job in selected[:10]:
        log(f"  {job['new'] or '  '} [{job['score']:>2}] {job['budget_text']:<16} {job['title'][:40]}")
        log(f"          {job['url']}")
    log("")
    log(f"一覧: {csv_path}")
    log(f"除外分: {excluded_path}")
    if wrote_xlsx:
        log(f"Excel: {xlsx_path}")
    return selected, excluded
